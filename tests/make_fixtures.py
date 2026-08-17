"""
Generate the test fixtures.

    python -m tests.make_fixtures          # create anything missing
    python -m tests.make_fixtures --force  # rebuild everything

The fixtures are committed, so this is only needed if one is deleted or you want
to change them. Everything is **generated rather than recorded**, which is what
makes the voice test a real round trip: gTTS produces the Hindi audio, Whisper
transcribes it back, and the two are compared. Nothing to download, nothing to
record, and the test cannot drift from a clip nobody can regenerate.

Needs a network connection (gTTS calls Google) — hence committing the output.
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

FIXTURES = Path(__file__).resolve().parent / "fixtures"

# The claim used throughout: it is one the corpus can genuinely verify, so the
# fixtures exercise the *verified* path rather than abstention.
CLAIM_EN = "India's software exports reached 222 billion dollars in 2024-25"
CLAIM_HI = "भारत का सॉफ्टवेयर निर्यात 2024-25 में 222 अरब डॉलर तक पहुंच गया"


def _audio(force: bool) -> None:
    from gtts import gTTS

    for name, text, lang in [("claim_hi.mp3", CLAIM_HI, "hi"),
                             ("claim_en.mp3", CLAIM_EN, "en")]:
        path = FIXTURES / name
        if path.exists() and not force:
            print(f"  exists  {name}")
            continue
        gTTS(text=text, lang=lang).save(str(path))
        print(f"  wrote   {name}  ({path.stat().st_size:,} bytes)")


def _images(force: bool) -> None:
    """
    Fake social-media screenshots.

    Deliberately include the interface furniture — like counts, "View all 302
    comments", a timestamp — because stripping that is what
    `adapt._tidy_ocr` exists to do, and a clean text image would not test it.
    """
    from PIL import Image, ImageDraw, ImageFont

    try:
        big = ImageFont.truetype("arial.ttf", 30)
        small = ImageFont.truetype("arial.ttf", 20)
    except OSError:
        big = small = ImageFont.load_default()

    specs = [
        ("insta_post.png", "factsdaily_official",
         ["India's software exports reached 222 billion",
          "dollars in 2024-25."],
         ["12,431 likes", "View all 302 comments", "2 days ago"]),
        # Text that is NOT a claim — should be rejected at the gate, which is
        # the answer the image workflow actually needs.
        ("insta_opinion.png", "randomuser",
         ["what time does the match start tonight?"],
         ["88 likes", "View all 12 comments"]),
    ]

    for name, handle, body, chrome in specs:
        path = FIXTURES / name
        if path.exists() and not force:
            print(f"  exists  {name}")
            continue
        height = 160 + 44 * len(body) + 34 * len(chrome)
        image = Image.new("RGB", (900, height), "white")
        draw = ImageDraw.Draw(image)
        draw.text((40, 28), handle, fill="black", font=small)
        y = 88
        for line in body:
            draw.text((40, y), line, fill="black", font=big)
            y += 44
        y += 30
        for line in chrome:
            draw.text((40, y), line, fill="#888", font=small)
            y += 34
        image.save(path)
        print(f"  wrote   {name}  ({path.stat().st_size:,} bytes)")


def _documents(force: bool) -> None:
    """
    A document asserting something the corpus contradicts.

    This is the fixture for the safety property: attaching it *and* claiming
    what it says must not yield TRUE. The corpus says 222 billion; this says 40.
    """
    import docx

    path = FIXTURES / "false_report.docx"
    if path.exists() and not force:
        print("  exists  false_report.docx")
        return

    document = docx.Document()
    document.add_heading("Internal Market Briefing", 0)
    document.add_paragraph(
        "Our analysis confirms that India's software exports collapsed to only "
        "40 billion dollars in 2024-25, a catastrophic decline of over 80 "
        "percent from the previous year. This figure has been verified by our "
        "research desk."
    )
    document.add_paragraph(
        "The decline was driven by a sharp contraction in demand across all "
        "major markets, according to figures compiled internally."
    )
    document.save(str(path))
    print(f"  wrote   false_report.docx  ({path.stat().st_size:,} bytes)")


def main() -> int:
    parser = argparse.ArgumentParser(prog="tests.make_fixtures")
    parser.add_argument("--force", action="store_true",
                        help="rebuild fixtures that already exist")
    args = parser.parse_args()

    FIXTURES.mkdir(parents=True, exist_ok=True)
    print(f"fixtures -> {FIXTURES}")

    for label, builder in (("audio (needs network)", _audio),
                           ("images", _images),
                           ("documents", _documents)):
        print(f"\n{label}")
        try:
            builder(args.force)
        except Exception as exc:
            print(f"  FAILED: {type(exc).__name__}: {exc}", file=sys.stderr)

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
