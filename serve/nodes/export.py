"""
Export a finished report as HTML, DOCX, or spoken audio.

Kept separate from `render.py`, which owns the markdown. Markdown is the
canonical form — everything here is a transformation of the same
`VerificationReport`, never an independent parse of the prose. That separation
is the fix for a real bug in the old system: the exporter regex-parsed the
report text to find the verdict, so the downloaded file could state a different
verdict from the API response for the same run.

Exports are lazy. Nothing here runs unless something asks for the format, so a
CLI verification never pays for DOCX generation and the audio path never touches
the network unless the input was actually voice.
"""

from __future__ import annotations

import logging
import re
from datetime import datetime, timezone
from pathlib import Path

from core.config import cfg
from serve.schemas import VerificationReport, VerifyState

logger = logging.getLogger(__name__)

# Verdict → (label, colour). Colour is used by the HTML export only.
_VERDICT_STYLE = {
    "TRUE":         ("TRUE", "#0f7b3f"),
    "MOSTLY_TRUE":  ("MOSTLY TRUE", "#4a8c2a"),
    "UNVERIFIABLE": ("UNVERIFIABLE", "#8a6d1f"),
    "MOSTLY_FALSE": ("MOSTLY FALSE", "#a8500f"),
    "FALSE":        ("FALSE", "#a11b1b"),
}


def _slug(text: str, limit: int = 40) -> str:
    """Filesystem-safe stem from a claim, in any script."""
    cleaned = re.sub(r"[^\w\s-]", "", text, flags=re.UNICODE).strip()
    cleaned = re.sub(r"[\s_]+", "-", cleaned)
    return (cleaned[:limit] or "report").lower()


def _output_path(state: VerifyState, suffix: str) -> Path:
    cfg.output_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M%S")
    stem = _slug(state.get("claim", "report"))
    return cfg.output_dir / f"{stamp}-{stem}{suffix}"


# ── HTML ─────────────────────────────────────────────────────────────────────

def to_html(state: VerifyState) -> Path:
    """
    Self-contained HTML — inline CSS, no external requests.

    Deliberately not a markdown-to-HTML conversion: the structured report is
    right there, so the verdict gets a real colour treatment and the evidence
    gets a table, neither of which survives a generic converter.
    """
    report: VerificationReport | None = state.get("report")
    documents = state.get("evidence") or []
    language = state.get("language", "en")

    def esc(value) -> str:
        return (str(value).replace("&", "&amp;").replace("<", "&lt;")
                .replace(">", "&gt;"))

    if report is None:
        body = (f"<h1>Not a verifiable claim</h1>"
                f"<blockquote>{esc(state.get('claim', ''))}</blockquote>"
                f"<p>{esc(state.get('reject_reason', ''))}</p>")
    else:
        label, colour = _VERDICT_STYLE.get(report.verdict,
                                           (report.verdict, "#555"))
        rows = []
        for assessment in report.evidence_analysis:
            index = assessment.evidence_index
            meta = (documents[index - 1].metadata
                    if 1 <= index <= len(documents) else {})
            flag = ""
            if meta.get("corpus_tier") == "user_upload":
                flag = " <em>(your upload — unverified)</em>"
            elif meta.get("corpus_tier") == "web":
                flag = " <em>(web search)</em>"
            rows.append(
                f"<tr><td>{index}</td>"
                f"<td>{esc(assessment.source)}{flag}</td>"
                f"<td>{meta.get('credibility', 0.0):.2f}</td>"
                f"<td>{esc(assessment.stance)}</td>"
                f"<td>{esc(assessment.reasoning)}</td></tr>"
            )

        findings = "".join(
            f"<li>{esc(f.finding)} <span class='cite'>[{f.evidence_index}]</span></li>"
            for f in report.key_findings
        )
        limitations = "".join(f"<li>{esc(l)}</li>" for l in report.limitations)

        body = f"""
        <h1>{esc(report.claim)}</h1>
        <div class="verdict" style="border-color:{colour}">
          <span class="label" style="color:{colour}">{label}</span>
          <span class="conf">{report.confidence}% confidence</span>
        </div>
        {'<h2>Key findings</h2><ul>' + findings + '</ul>' if findings else ''}
        <h2>Evidence</h2>
        <table><thead><tr><th>#</th><th>Source</th><th>Cred.</th>
        <th>Stance</th><th>Assessment</th></tr></thead>
        <tbody>{''.join(rows)}</tbody></table>
        {'<h2>Limitations</h2><ul>' + limitations + '</ul>' if limitations else ''}
        <h2>Conclusion</h2><p>{esc(report.conclusion)}</p>
        """

    direction = "rtl" if language in {"ar", "ur", "fa", "he"} else "ltr"
    html = f"""<!doctype html>
<html lang="{esc(language)}" dir="{direction}"><head>
<meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>VerifAI — {esc(state.get('claim', '')[:60])}</title>
<style>
  :root {{ color-scheme: light dark; }}
  body {{ font-family: system-ui,-apple-system,"Segoe UI",Roboto,sans-serif;
         max-width: 46rem; margin: 2rem auto; padding: 0 1.25rem;
         line-height: 1.6; }}
  h1 {{ font-size: 1.5rem; line-height: 1.3; }}
  h2 {{ font-size: 1.05rem; margin-top: 2rem; text-transform: uppercase;
        letter-spacing: .06em; opacity: .7; }}
  .verdict {{ border-left: 5px solid; padding: .75rem 1rem; margin: 1.25rem 0;
             background: rgba(127,127,127,.08); }}
  .label {{ font-weight: 700; font-size: 1.2rem; }}
  .conf {{ opacity: .7; margin-inline-start: .75rem; }}
  table {{ border-collapse: collapse; width: 100%; font-size: .9rem; }}
  th,td {{ text-align: start; padding: .5rem; border-bottom: 1px solid rgba(127,127,127,.3);
           vertical-align: top; }}
  .cite {{ opacity: .6; font-size: .85em; }}
  footer {{ margin-top: 3rem; font-size: .8rem; opacity: .6; }}
</style></head><body>
{body}
<footer>Generated by VerifAI · outcome: {esc(state.get('outcome', ''))} ·
{len(documents)} evidence item(s) ·
{state.get('diagnostics', {}).get('llm_calls', 0)} LLM call(s)</footer>
</body></html>"""

    path = _output_path(state, ".html")
    path.write_text(html, encoding="utf-8")
    logger.info("Wrote %s", path.name)
    return path


# ── DOCX ─────────────────────────────────────────────────────────────────────

def to_docx(state: VerifyState) -> Path:
    """Word document. Optional dependency — raises a clear message if absent."""
    try:
        import docx
    except ImportError:
        raise RuntimeError(
            "DOCX export needs python-docx: pip install -e '.[docs]'"
        ) from None

    report: VerificationReport | None = state.get("report")
    documents = state.get("evidence") or []

    document = docx.Document()
    document.add_heading("Fact verification report", level=0)

    if report is None:
        document.add_paragraph(state.get("claim", ""), style="Intense Quote")
        document.add_paragraph(state.get("reject_reason", ""))
    else:
        document.add_paragraph(report.claim, style="Intense Quote")

        label = _VERDICT_STYLE.get(report.verdict, (report.verdict, ""))[0]
        verdict = document.add_paragraph()
        verdict.add_run(f"{label}  ").bold = True
        verdict.add_run(f"({report.confidence}% confidence)")

        if report.key_findings:
            document.add_heading("Key findings", level=1)
            for finding in report.key_findings:
                document.add_paragraph(
                    f"{finding.finding}  [{finding.evidence_index}]",
                    style="List Bullet",
                )

        if report.evidence_analysis:
            document.add_heading("Evidence", level=1)
            table = document.add_table(rows=1, cols=4)
            table.style = "Light Grid Accent 1"
            for cell, text in zip(table.rows[0].cells,
                                  ("#", "Source", "Credibility", "Stance")):
                cell.text = text
            for assessment in report.evidence_analysis:
                index = assessment.evidence_index
                meta = (documents[index - 1].metadata
                        if 1 <= index <= len(documents) else {})
                cells = table.add_row().cells
                cells[0].text = str(index)
                cells[1].text = assessment.source + (
                    " (your upload — unverified)"
                    if meta.get("corpus_tier") == "user_upload" else ""
                )
                cells[2].text = f"{meta.get('credibility', 0.0):.2f}"
                cells[3].text = assessment.stance

        if report.limitations:
            document.add_heading("Limitations", level=1)
            for limitation in report.limitations:
                document.add_paragraph(limitation, style="List Bullet")

        document.add_heading("Conclusion", level=1)
        document.add_paragraph(report.conclusion)

    path = _output_path(state, ".docx")
    document.save(str(path))
    logger.info("Wrote %s", path.name)
    return path


# ── Audio ────────────────────────────────────────────────────────────────────

# gTTS language codes that differ from our ISO codes, plus the ones it cannot do.
_GTTS_LANG = {"zh": "zh-CN"}
_GTTS_UNSUPPORTED = {"pa"}   # Punjabi has no gTTS voice


def to_audio(state: VerifyState) -> Path | None:
    """
    Speak the verdict and conclusion back, in the report's own language.

    Only the verdict line and conclusion — a spoken evidence table is unusable,
    and the point of audio output is that someone who sent a voice note gets an
    answer they can listen to.

    Returns None rather than raising when the language has no voice: losing the
    audio is not a reason to fail a verification that otherwise succeeded.
    """
    report: VerificationReport | None = state.get("report")
    if report is None:
        return None

    language = state.get("language", "en")
    if language in _GTTS_UNSUPPORTED:
        logger.info("No gTTS voice for %r — skipping audio", language)
        return None

    try:
        from gtts import gTTS
    except ImportError:
        logger.info("gTTS not installed — skipping audio")
        return None

    label = _VERDICT_STYLE.get(report.verdict, (report.verdict, ""))[0]
    spoken = f"{label}. {report.confidence} percent confidence. {report.conclusion}"

    try:
        # gTTS reaches translate.google.com. That host must be in
        # VERIFAI_DNS_PREWARM or the lookup crashes the process — see
        # core.init_native_libs.
        tts = gTTS(text=spoken, lang=_GTTS_LANG.get(language, language))
        path = _output_path(state, ".mp3")
        tts.save(str(path))
        logger.info("Wrote %s", path.name)
        return path
    except Exception as exc:
        logger.warning("Audio generation failed (%s) — continuing without it", exc)
        return None


# ── Dispatcher ───────────────────────────────────────────────────────────────

def export(state: VerifyState, formats: list[str]) -> dict[str, str]:
    """Render the requested formats, skipping any that fail."""
    handlers = {"html": to_html, "docx": to_docx, "mp3": to_audio}
    artifacts: dict[str, str] = {}

    for name in formats:
        handler = handlers.get(name)
        if handler is None:
            logger.warning("Unknown export format %r", name)
            continue
        try:
            path = handler(state)
            if path is not None:
                artifacts[name] = str(path)
        except Exception as exc:
            logger.warning("%s export failed: %s", name, exc)

    return artifacts
